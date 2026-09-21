"""Excel report export — mirrors modCostbookReport 4+4 sheet layouts."""
from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.services.costbook_engine import CostbookEngine

NAVY = "1F2F69"
PARETO_FILL = PatternFill(start_color="92D050", end_color="92D050", fill_type="solid")


def _brand_sheet(ws, title: str, subtitle: str) -> None:
    ws["B1"] = title
    ws["B1"].font = Font(name="Segoe UI", size=14, bold=True, color=NAVY)
    ws["B2"] = subtitle
    ws["B2"].font = Font(name="Segoe UI", size=9, italic=True, color="969696")


def export_costbook_report(engine: CostbookEngine, costbook: dict[str, Any]) -> bytes:
    comp = engine.calculate_comparison([costbook])
    wb = Workbook()
    ws = wb.active
    ws.title = "Overview"
    _brand_sheet(ws, "Costbook Report — Overview", costbook.get("display", costbook["vehicle_code"]))
    r = 4
    for k, v in engine.carline_fields(costbook["vehicle_code"]).items():
        ws.cell(r, 2, k)
        ws.cell(r, 3, v)
        r += 1
    r += 1
    ws.cell(r, 2, "Total TPC")
    ws.cell(r, 3, comp["kpis"]["baseline_tpc"])
    r += 2
    ws.cell(r, 2, "Top Systems")
    r += 1
    for item in comp["systems"][:10]:
        ws.cell(r, 2, item["label"])
        ws.cell(r, 3, item["tpc"])
        ws.cell(r, 4, item["pct"])
        r += 1

    ws2 = wb.create_sheet("Full BOM")
    _brand_sheet(ws2, "Full BOM", "Pareto-ranked lines")
    headers = ["Part Number", "Description", "Macro System", "TPC", "TPC %", "Cum %"]
    for c, h in enumerate(headers, 2):
        ws2.cell(4, c, h).font = Font(bold=True, color=NAVY)
    for i, row in enumerate(comp["full_bom"], 5):
        ws2.cell(i, 2, row["part_number"])
        ws2.cell(i, 3, row["part_description"])
        ws2.cell(i, 4, row["macro_system"])
        ws2.cell(i, 5, row["tpc"])
        ws2.cell(i, 6, row["pct"])
        ws2.cell(i, 7, row["cumulative_pct"])
        if row["cumulative_pct"] <= 0.8:
            ws2.cell(i, 5).fill = PARETO_FILL

    ws3 = wb.create_sheet("Hierarchical View")
    _brand_sheet(ws3, "Hierarchical View", "L1 → L2 → L3 → Part")
    ws3.cell(4, 2, "Path").font = Font(bold=True)
    ws3.cell(4, 3, "TPC").font = Font(bold=True)
    for i, row in enumerate(comp["hierarchical"], 5):
        ws3.cell(i, 2, " > ".join(row["path"]))
        ws3.cell(i, 3, row["tpc"])

    ws4 = wb.create_sheet("Hierarchical by 5th")
    _brand_sheet(ws4, "Hierarchical View by 5th", "")
    ws4.cell(4, 2, "Path").font = Font(bold=True)
    ws4.cell(4, 3, "TPC").font = Font(bold=True)
    for i, row in enumerate(comp["hierarchical_by_fifth"], 5):
        ws4.cell(i, 2, " > ".join(row["path"]))
        ws4.cell(i, 3, row["tpc"])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_gap_report(engine: CostbookEngine, costbooks: list[dict[str, Any]]) -> bytes:
    comp = engine.calculate_comparison(costbooks)
    wb = Workbook()
    ws = wb.active
    ws.title = "Comparison Overview"
    _brand_sheet(ws, "Gap Comparison — Overview", f"{len(costbooks)} costbooks")
    ws.cell(4, 2, "Vehicle").font = Font(bold=True)
    ws.cell(4, 3, "Milestone").font = Font(bold=True)
    ws.cell(4, 4, "Date").font = Font(bold=True)
    ws.cell(4, 5, "TPC").font = Font(bold=True)
    for i, cb in enumerate(comp["costbooks"], 5):
        ws.cell(i, 2, cb["vehicle_code"])
        ws.cell(i, 3, cb["milestone"])
        ws.cell(i, 4, cb["milestone_date"])
        ws.cell(i, 5, cb["total_tpc"])

    ws2 = wb.create_sheet("Full BOM Comparison")
    _brand_sheet(ws2, "Full BOM Comparison", "Aligned by part number")
    ws2.cell(4, 2, "Part Number").font = Font(bold=True)
    ws2.cell(4, 3, "Description").font = Font(bold=True)
    for vi in range(len(costbooks)):
        ws2.cell(4, 4 + vi, f"V{vi + 1} TPC").font = Font(bold=True)
    ws2.cell(4, 4 + len(costbooks), "Gap").font = Font(bold=True)
    for i, row in enumerate(comp["aligned_bom"][:500], 5):
        ws2.cell(i, 2, row.get("Part Number", ""))
        ws2.cell(i, 3, row.get("Part Description", ""))
        for vi, val in enumerate(row["values"]):
            ws2.cell(i, 4 + vi, val)
        ws2.cell(i, 4 + len(costbooks), row["gap"])

    ws3 = wb.create_sheet("Hierarchical Gaps")
    _brand_sheet(ws3, "Hierarchical Gaps", comp["level"])
    ws3.cell(4, 2, "Label").font = Font(bold=True)
    for vi in range(len(costbooks)):
        ws3.cell(4, 3 + vi, f"V{vi + 1}").font = Font(bold=True)
    ws3.cell(4, 3 + len(costbooks), "Gap").font = Font(bold=True)
    for i, row in enumerate(comp["gap_drivers"][:200], 5):
        ws3.cell(i, 2, row["label"])
        for vi, val in enumerate(row["values"]):
            ws3.cell(i, 3 + vi, val)
        ws3.cell(i, 3 + len(costbooks), row["gap"])

    ws4 = wb.create_sheet("Hierarchical Gaps by 5th")
    _brand_sheet(ws4, "Hierarchical Gaps by 5th", "")
    gap5 = engine.gap_agg(costbooks, "5th")
    ws4.cell(4, 2, "5th").font = Font(bold=True)
    for vi in range(len(costbooks)):
        ws4.cell(4, 3 + vi, f"V{vi + 1}").font = Font(bold=True)
    ws4.cell(4, 3 + len(costbooks), "Gap").font = Font(bold=True)
    for i, row in enumerate(gap5[:200], 5):
        ws4.cell(i, 2, row["label"])
        for vi, val in enumerate(row["values"]):
            ws4.cell(i, 3 + vi, val)
        ws4.cell(i, 3 + len(costbooks), row["gap"])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
